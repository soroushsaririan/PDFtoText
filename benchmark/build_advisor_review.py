"""Build a single reviewable file for an advisor: original crop image,
printed field label, prompt style, Qwen's actual output, ground truth, and
whether it matched -- plus a sheet explaining how accuracy is measured.

    ./.venv/bin/python benchmark/build_advisor_review.py

Reads the two already-scored sheets (answer style and discriminative style,
both Qwen2.5-VL-7B-Instruct-4bit via MLX -- the best-performing config) and
writes one .xlsx with the crop image embedded directly in each row.
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

BENCH_DIR = Path(__file__).resolve().parent
OUT_PATH = Path.home() / "Downloads" / "Qwen Handwriting Results - Advisor Review.xlsx"

SOURCES = [
    ("Answer style", BENCH_DIR / "real_ground_truth_sample_scored_mlx_answer_none.xlsx"),
    ("Discriminative style", BENCH_DIR / "real_ground_truth_sample_scored_mlx_discriminative_none.xlsx"),
]

HEADER_FILL = PatternFill(start_color="1F2A44", end_color="1F2A44", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MATCH_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
NOMATCH_FILL = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
ROW_HEIGHT = 70
IMG_TARGET_H = 60   # native crop height is ~40-65px -- keep thumbnails close to
IMG_MAX_W = 560     # native size rather than over-compressing wide crops down to a sliver
SAMPLE_SIZE = 5     # one field per participant, same fields on both style sheets


def add_methodology_sheet(wb, style_accuracy):
    ws = wb.active
    ws.title = "Methodology"
    ws.column_dimensions["A"].width = 100
    results_lines = [
        (f"{title}: {correct} of {total} correct ({correct / total:.0%}).", False)
        for title, correct, total in style_accuracy
    ]
    lines = [
        ("Model and setup", True),
        ("Qwen2.5-VL-7B-Instruct-4bit, run locally via Apple's MLX framework "
         "(4-bit quantized so it fits Apple Silicon unified memory). This is the "
         "best-performing configuration found so far.", False),
        ("", False),
        ("Ground truth", True),
        ("44 individual form fields (numbers, pain scores, short text) were "
         "cropped from real, de-identified study forms across 5 participants "
         "(VPT001, VPT002, VPT016, VPT018, VPT020), chosen to cover known data "
         "irregularities. A human read each crop directly from the original PDF "
         "and typed the correct value -- this is the \"ground truth\" column on "
         "each results sheet. The accuracy percentages below are computed over "
         "all 44; this file shows 5 representative fields (one per participant) "
         "for a readable walkthrough, not the full set.", False),
        ("", False),
        ("What the model was shown", True),
        ("For each field, the model was shown ONLY the cropped image of that "
         "field (not the whole form) plus a text prompt naming the printed "
         "question for that field (the \"label\" column) and, for the "
         "discriminative style, a structured set of sub-questions to answer "
         "about it. The model was never shown the ground truth.", False),
        ("", False),
        ("Two prompt styles compared", True),
        ("\"Answer style\" asks for the bare value only (e.g. just \"7\"). "
         "\"Discriminative style\" asks the model to fill in several numbered "
         "sub-fields per crop (is it legible, was anything crossed out, the "
         "value itself, and a one-line justification of what strokes it's "
         "reading) -- the \"Qwen answer (extracted)\" column on that sheet is "
         "just the VALUE sub-field pulled out of that longer structured reply "
         "for comparison; the full reply is also kept in \"Qwen raw output\".", False),
        ("", False),
        ("How accuracy is scored", True),
        ("The model's answer is compared to ground truth as an exact string "
         "match after light normalization: lowercased, extra whitespace "
         "collapsed, surrounding quotes stripped, and \"4.0\" treated the same "
         "as \"4\". Answers meaning \"nothing was legible\" (e.g. \"[illegible]\") "
         "are treated as equivalent to each other, and likewise for \"the field "
         "was blank\". No partial credit -- a value that is off by one digit, "
         "or a text answer with one wrong character, counts as incorrect. This "
         "is a strict measure; it likely understates how close a wrong answer "
         "sometimes was.", False),
        ("", False),
        ("Results", True),
        *results_lines,
        ("See the two results sheets for the field-by-field detail behind "
         "those numbers, including the original crop image for each field.", False),
    ]
    for i, (text, is_header) in enumerate(lines, start=1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if is_header:
            cell.font = Font(bold=True, size=13)
        ws.row_dimensions[i].height = 18 if not text else (22 if is_header else 45)
    return ws


def add_results_sheet(wb, title, df):
    ws = wb.create_sheet(title)
    headers = ["Original crop", "Participant", "Field label", "Field kind",
               "Qwen raw output", "Qwen answer (extracted)", "Ground truth", "Match?"]
    widths = [80, 12, 30, 10, 40, 20, 16, 8]
    for col, (h, w) in enumerate(zip(headers, widths), start=1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A2"

    for i, row in enumerate(df.itertuples(), start=2):
        ws.row_dimensions[i].height = ROW_HEIGHT
        crop_path = BENCH_DIR / row.crop
        if crop_path.exists():
            with PILImage.open(crop_path) as im:
                w, h = im.size
                scale = IMG_TARGET_H / h
                new_w, new_h = max(1, int(w * scale)), IMG_TARGET_H
                if new_w > IMG_MAX_W:
                    scale = IMG_MAX_W / w
                    new_w, new_h = IMG_MAX_W, max(1, int(h * scale))
                im = im.convert("RGB").resize((new_w, new_h))
                tmp = crop_path.with_name(f"_thumb_{crop_path.stem}.png")
                im.save(tmp)
            xl_img = XLImage(str(tmp))
            ws.add_image(xl_img, f"A{i}")
        ws.cell(row=i, column=2, value=row.participant)
        ws.cell(row=i, column=3, value=row.label)
        ws.cell(row=i, column=4, value=row.field_kind)
        raw_cell = ws.cell(row=i, column=5, value=str(row.raw_result))
        raw_cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=i, column=6, value=str(row.result))
        ws.cell(row=i, column=7, value=str(row.ground_truth))
        match_cell = ws.cell(row=i, column=8, value="Yes" if row.match else "No")
        match_cell.fill = MATCH_FILL if row.match else NOMATCH_FILL
    return ws


def pick_sample_crops(df, n):
    """One field per participant, same crops used on every style sheet so
    rows compare directly. A sample this small must still be a fair
    representation of the real accuracy: picking, say, "participant_id" for
    every participant would show 5 wrong answers in a row (it's one of the
    weaker fields) and badly misrepresent a model that's actually right 61%
    of the time overall. So: target a correct/incorrect split proportional
    to the real accuracy, and prefer a field_kind not already shown."""
    overall_rate = df["match"].mean()
    target_correct = round(n * overall_rate)
    participants = list(dict.fromkeys(df["participant"]))[:n]

    chosen, used_kinds, correct_so_far = [], set(), 0
    for i, participant in enumerate(participants):
        group = df[df["participant"] == participant]
        slots_left = n - i
        need_correct_now = (target_correct - correct_so_far) >= slots_left
        pool = group[group["match"]] if need_correct_now else group[~group["match"]]
        if pool.empty:
            pool = group
        unseen_kind = pool[~pool["field_kind"].isin(used_kinds)]
        row = (unseen_kind if not unseen_kind.empty else pool).iloc[0]
        chosen.append(row["crop"])
        used_kinds.add(row["field_kind"])
        correct_so_far += int(row["match"])
    return chosen


def main():
    for _, path in SOURCES:
        if not path.exists():
            print(f"Missing: {path}")
            sys.exit(1)

    full_dfs = [pd.read_excel(path) for _, path in SOURCES]
    style_accuracy = [(title, int(df["match"].sum()), len(df))
                       for (title, _), df in zip(SOURCES, full_dfs)]

    wb = Workbook()
    add_methodology_sheet(wb, style_accuracy)

    sample_crops = pick_sample_crops(full_dfs[0], SAMPLE_SIZE)
    for (title, _), df in zip(SOURCES, full_dfs):
        df = df[df["crop"].isin(sample_crops)].reset_index(drop=True)
        add_results_sheet(wb, title, df)
        print(f"Added sheet '{title}': {len(df)} rows")

    wb.save(OUT_PATH)
    # clean up temp thumbnails
    for f in BENCH_DIR.glob("**/_thumb_*.png"):
        f.unlink()
    print(f"\nSaved: {OUT_PATH}")


if __name__ == "__main__":
    main()
