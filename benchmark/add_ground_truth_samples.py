"""Cut answer crops for more real participants and append them to
real_ground_truth_sample.xlsx with ground_truth left blank, for a person
to fill in from the images.

    ./.venv/bin/python benchmark/add_ground_truth_samples.py \\
        --base-dir "~/Desktop/TextAnalysis VPT files backup 20260820" --batch 2

By default it takes every participant under --base-dir that isn't in the
sheet yet. New rows follow the sheet's existing layout: participants in
descending order; within a participant, weekly session forms (latest week
first), then the final CRF visit, then visit-1 (intake) CRFs; within a
form, the fields in FIELDS_BY_FORM order. participant_id's ground truth is
prefilled from the folder name, as run_prompt_styles.py does.

Crops go to real_ground_truth_sample_crops_<batch>/ (page renders under
its pages/). Next to each crop is <name>_context.png: the same field with
the printed question and a couple of rows above and below, because the
answer crop on its own can clip the handwriting (people write above the
underline) -- a "7" cut off at the top looks like a "1".
"""

import argparse
import sys
from pathlib import Path

import fitz
import openpyxl

BENCH_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BENCH_DIR.parent / "src"))
sys.path.insert(0, str(BENCH_DIR))

import prompts  # noqa: E402
from field_extraction import FINAL_LABELS, INTAKE_LABELS, WEEKLY_LABELS  # noqa: E402
from ocr_pipeline import UnreadableFile  # noqa: E402
from field_reader import (  # noqa: E402
    answer_crop,
    field_kind_for,
    label_hits,
    line_index,
    next_label_start_x,
    open_page,
    page_words,
)
from ocr_pipeline import render_file_to_images  # noqa: E402
from run_prompt_styles import resolve_jobs  # noqa: E402

SHEET_PATH = BENCH_DIR / "real_ground_truth_sample.xlsx"

# The fields batches 1 and 2 sampled. These were the only labels that
# actually matched the printed forms at the time, so they are the only
# fields with any measured accuracy; --fields all (or an explicit list)
# samples the rest, which have never been scored against ground truth.
FIELDS_BY_FORM = {
    "weekly": ("weekly_visit", "participant_id", "pain_now", "mileage"),
    "final": ("pain_now",),
    "intake": ("injury_location", "tester", "participant_id", "age"),
}
LABEL_SETS = {"weekly": WEEKLY_LABELS, "final": FINAL_LABELS, "intake": INTAKE_LABELS}

CONTEXT_ROWS_ABOVE = 2.5
CONTEXT_ROWS_BELOW = 1.5


def _context_crop(page_image, label_box, max_x1):
    x0, y0, x1, y1 = label_box
    line_h = max(y1 - y0, 12)
    right = max_x1 if max_x1 is not None else page_image.width
    return page_image.crop((
        max(0, x0 - 10), max(0, round(y0 - line_h * CONTEXT_ROWS_ABOVE)),
        min(page_image.width, right + 10), min(page_image.height, round(y1 + line_h * CONTEXT_ROWS_BELOW)),
    ))


def collect(participant, pdf_path, form, week, crops_dir, fields_by_form=None):
    wanted = (fields_by_form or FIELDS_BY_FORM)[form]
    label_set = [t for t in LABEL_SETS[form] if t[1] in wanted]
    pages = render_file_to_images(pdf_path, crops_dir / "pages" / pdf_path.stem)
    found = []
    for page_no, img_path in enumerate(pages, start=1):
        page_image = open_page(img_path)
        data = page_words(page_image)
        lines = line_index(data)
        for label, key, kind in label_set:
            field_kind = field_kind_for(key, kind, label)
            for occ, (box, line_key) in enumerate(label_hits(data, lines, label), start=1):
                max_x1 = next_label_start_x(data, lines.get(line_key, []), box[2])
                crop = answer_crop(page_image, box, field_kind, max_x1=max_x1)
                if crop is None:
                    continue
                wk_tag = f"_wk{week}" if week is not None else ""
                name = f"{pdf_path.stem}{wk_tag}_p{page_no:02d}_{key}_{occ}"
                crop.save(crops_dir / f"{name}.png")
                _context_crop(page_image, box, max_x1).save(crops_dir / f"{name}_context.png")
                found.append({
                    "participant": participant,
                    "pdf": pdf_path.name,
                    "form_type": form,
                    "week": week,
                    "page": page_no,
                    "field_key": key,
                    "label": label,
                    "field_kind": field_kind,
                    "occurrence": occ,
                    "style": "answer",
                    "ground_truth": participant if key == "participant_id" else None,
                    "result": None,
                    "crop": f"{crops_dir.name}/{name}.png",
                    "prompt": prompts.build_prompt(label_context=label, field_kind=field_kind, style="answer"),
                })
    found.sort(key=lambda r: (wanted.index(r["field_key"]), r["page"], r["occurrence"]))
    return found


def participant_rows(participant_dir, crops_dir, fields_by_form=None):
    jobs = resolve_jobs(participant_dir, ["weekly", "intake", "final"])
    weekly = sorted([j for j in jobs if j[0] == "weekly"], key=lambda j: -j[3])
    final = [j for j in jobs if j[0] == "final"]
    intake = [j for j in jobs if j[0] == "intake"]
    rows = []
    for form, pdf_path, _, week in weekly + final + intake:
        print(f"  {participant_dir.name}: {pdf_path.name}", flush=True)
        try:
            rows.extend(collect(participant_dir.name, pdf_path, form, week, crops_dir, fields_by_form))
        except (UnreadableFile, fitz.FileDataError) as exc:
            # seen: an iPhone HEIF photo saved with a .pdf extension
            print(f"    SKIPPED, can't render: {exc}", flush=True)
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--base-dir", required=True, help="Folder with the real VPT001, VPT002, ... subfolders.")
    parser.add_argument("--batch", required=True, help="Batch name for the crops folder, e.g. 2.")
    parser.add_argument("--fields", nargs="*", default=None,
                        help="Field keys to sample, 'all' for every label on each form "
                             "(default: the keys batches 1-2 used).")
    parser.add_argument("--participants", nargs="*", default=None,
                        help="Only these participants (default: every one not already in the sheet).")
    args = parser.parse_args()

    base_dir = Path(args.base_dir).expanduser()
    wb = openpyxl.load_workbook(SHEET_PATH)
    ws = wb.active
    header = [c.value for c in ws[1]]
    already = {ws.cell(r, header.index("participant") + 1).value for r in range(2, ws.max_row + 1)}

    participants = args.participants or sorted(
        p.name for p in base_dir.iterdir() if p.is_dir() and p.name.startswith("VPT") and p.name not in already
    )
    if args.fields == ["all"]:
        fields_by_form = {form: tuple(key for _, key, _ in labels) for form, labels in LABEL_SETS.items()}
    elif args.fields:
        fields_by_form = {form: tuple(key for _, key, _ in labels if key in set(args.fields))
                          for form, labels in LABEL_SETS.items()}
    else:
        fields_by_form = FIELDS_BY_FORM

    crops_dir = BENCH_DIR / f"real_ground_truth_sample_crops_{args.batch}"
    crops_dir.mkdir(exist_ok=True)

    rows = []
    for participant in sorted(participants, reverse=True):
        rows.extend(participant_rows(base_dir / participant, crops_dir, fields_by_form))

    for row in rows:
        ws.append([row.get(col) for col in header])
    wb.save(SHEET_PATH)
    print(f"\nAppended {len(rows)} rows for {len(participants)} participants to {SHEET_PATH.name}")
    print(f"Crops: {crops_dir}")


if __name__ == "__main__":
    main()
