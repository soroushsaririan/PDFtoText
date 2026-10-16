"""Type ground truth for the unlabelled crops in real_ground_truth_sample.xlsx.

    ./.venv/bin/python benchmark/label_sheet.py            # build the page
    open benchmark/label_sheet.html                        # fill it in
    ./.venv/bin/python benchmark/label_sheet.py --apply ~/Downloads/ground_truth.csv

Reading 356 crops out of Excel means scrolling a filename column and
opening each PNG by hand, which is slow enough that the fields sampled in
batch 3 have no measured accuracy at all. The page shows one field at a
time -- the printed question in context above, the answer crop enlarged
below -- and keeps what has been typed in the browser, so it can be closed
and reopened. Nothing leaves the machine: it is a local file pointing at
the crop folders on disk, and the real handwriting never goes near a
server.

--apply writes the exported CSV back into the sheet by row number, leaving
every other cell alone.
"""

import argparse
import csv
import json
from pathlib import Path

import openpyxl

BENCH_DIR = Path(__file__).resolve().parent
SHEET_PATH = BENCH_DIR / "real_ground_truth_sample.xlsx"
PAGE_PATH = BENCH_DIR / "label_sheet.html"


def unlabelled(sheet_path):
    """One entry per row whose ground_truth is still empty, in sheet order."""
    workbook = openpyxl.load_workbook(sheet_path)
    sheet = workbook.active
    header = [c.value for c in sheet[1]]
    column = {name: i + 1 for i, name in enumerate(header)}
    pending = []
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, column["ground_truth"]).value not in (None, ""):
            continue
        crop = sheet.cell(row, column["crop"]).value
        pending.append({
            "row": row,
            "participant": sheet.cell(row, column["participant"]).value,
            "week": sheet.cell(row, column["week"]).value,
            "form": sheet.cell(row, column["form_type"]).value,
            "field": sheet.cell(row, column["field_key"]).value,
            "kind": sheet.cell(row, column["field_kind"]).value,
            "label": sheet.cell(row, column["label"]).value,
            "crop": crop,
            "context": crop.replace(".png", "_context.png") if crop else None,
        })
    return pending


def write_page(pending, path=PAGE_PATH):
    path.write_text(PAGE_TEMPLATE.replace("__ITEMS__", json.dumps(pending)), encoding="utf-8")
    return path


def apply_csv(csv_path, sheet_path=SHEET_PATH):
    """Copy ground_truth values back into the sheet by row number.

    Blank values are skipped rather than written as empty strings, so a
    half-finished export can be applied and the page reopened for the rest.
    """
    workbook = openpyxl.load_workbook(sheet_path)
    sheet = workbook.active
    header = [c.value for c in sheet[1]]
    gt_column = header.index("ground_truth") + 1
    written = 0
    with open(csv_path, newline="", encoding="utf-8-sig") as handle:
        for entry in csv.DictReader(handle):
            value = (entry.get("ground_truth") or "").strip()
            if not value:
                continue
            sheet.cell(int(entry["row"]), gt_column).value = value
            written += 1
    workbook.save(sheet_path)
    return written


PAGE_TEMPLATE = r"""<!doctype html>
<meta charset="utf-8">
<title>Ground truth</title>
<style>
  :root { color-scheme: light dark; --edge: #8883; }
  body { font: 15px/1.5 system-ui, sans-serif; margin: 0; padding: 16px; max-width: 980px; }
  header { display: flex; gap: 12px; align-items: baseline; flex-wrap: wrap; }
  h1 { font-size: 17px; margin: 0; }
  .meta { color: #888; }
  .question { margin: 12px 0 4px; font-weight: 600; }
  figure { margin: 0 0 12px; }
  figcaption { color: #888; font-size: 13px; margin-bottom: 4px; }
  img { max-width: 100%; border: 1px solid var(--edge); border-radius: 4px; background: #fff; }
  #crop img { max-height: 190px; }
  form { display: flex; gap: 8px; flex-wrap: wrap; align-items: center; position: sticky; bottom: 0;
         padding: 12px 0; background: Canvas; border-top: 1px solid var(--edge); }
  input[type=text] { font: inherit; padding: 8px; min-width: 240px; }
  button { font: inherit; padding: 8px 12px; }
  progress { width: 220px; }
  .done { color: #2a7; }
</style>
<header>
  <h1>Ground truth</h1>
  <progress id="bar" value="0"></progress>
  <span class="meta" id="count"></span>
  <button id="export">Download CSV</button>
</header>
<p class="meta" id="where"></p>
<p class="question" id="question"></p>
<figure id="context"><figcaption>In context</figcaption></figure>
<figure id="crop"><figcaption>Answer crop</figcaption></figure>
<form id="entry" autocomplete="off">
  <input type="text" id="value" placeholder="what it says" autofocus>
  <button type="submit">Next &rarr;</button>
  <button type="button" id="blank">Blank</button>
  <button type="button" id="illegible">Illegible</button>
  <button type="button" id="back">&larr; Back</button>
  <span id="saved"></span>
</form>
<script>
const items = __ITEMS__;
const store = "ground-truth-" + items.length;
let answers = JSON.parse(localStorage.getItem(store) || "{}");
let at = items.findIndex(it => !(it.row in answers));
if (at < 0) at = items.length - 1;

const el = id => document.getElementById(id);

function show() {
  const it = items[at];
  el("where").textContent = `${it.participant} · ${it.form}${it.week ? " week " + it.week : ""} · ${it.field} (${it.kind})`;
  el("question").textContent = it.label;
  el("context").querySelector("img")?.remove();
  el("crop").querySelector("img")?.remove();
  for (const [slot, src] of [["context", it.context], ["crop", it.crop]]) {
    if (!src) continue;
    const img = document.createElement("img");
    img.src = src;
    el(slot).append(img);
  }
  el("value").value = answers[it.row] ?? "";
  const done = Object.keys(answers).length;
  el("bar").max = items.length;
  el("bar").value = done;
  el("count").textContent = `${done} of ${items.length} · showing ${at + 1}`;
  el("saved").textContent = it.row in answers ? "saved" : "";
  el("saved").className = it.row in answers ? "done" : "";
  el("value").focus();
}

function record(value) {
  answers[items[at].row] = value;
  localStorage.setItem(store, JSON.stringify(answers));
  at = Math.min(at + 1, items.length - 1);
  show();
}

el("entry").addEventListener("submit", e => { e.preventDefault(); record(el("value").value.trim()); });
el("blank").addEventListener("click", () => record("[blank]"));
el("illegible").addEventListener("click", () => record("[illegible]"));
el("back").addEventListener("click", () => { at = Math.max(at - 1, 0); show(); });
el("export").addEventListener("click", () => {
  const lines = ["row,ground_truth"];
  for (const it of items) {
    if (!(it.row in answers)) continue;
    lines.push(`${it.row},"${String(answers[it.row]).replace(/"/g, '""')}"`);
  }
  const url = URL.createObjectURL(new Blob([lines.join("\n")], {type: "text/csv"}));
  const a = document.createElement("a");
  a.href = url;
  a.download = "ground_truth.csv";
  a.click();
  URL.revokeObjectURL(url);
});
show();
</script>
"""


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--apply", metavar="CSV", help="Write an exported CSV back into the sheet.")
    parser.add_argument("--sheet", default=str(SHEET_PATH))
    args = parser.parse_args()

    if args.apply:
        written = apply_csv(args.apply, Path(args.sheet))
        print(f"Wrote {written} ground_truth value(s) into {Path(args.sheet).name}")
        return

    pending = unlabelled(Path(args.sheet))
    page = write_page(pending)
    by_field = {}
    for entry in pending:
        by_field[entry["field"]] = by_field.get(entry["field"], 0) + 1
    print(f"{len(pending)} unlabelled row(s) across {len(by_field)} field(s)")
    for field, count in sorted(by_field.items(), key=lambda kv: -kv[1]):
        print(f"  {field:24s} {count}")
    print(f"\nOpen: {page}")


if __name__ == "__main__":
    main()
